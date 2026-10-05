"""PPTX ingestion using python-pptx. One 'document' per slide, including
speaker notes, so 'explain slide 15' style queries can be answered precisely."""
from pptx import Presentation


def load_pptx(path: str) -> list[dict]:
    prs = Presentation(path)
    docs = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append("\t".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("[Speaker notes] " + slide.notes_slide.notes_text_frame.text.strip())

        if parts:
            docs.append({
                "text": "\n".join(parts),
                "location": f"slide {slide_num}",
                "page": slide_num,
            })
    return docs
