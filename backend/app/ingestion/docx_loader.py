"""DOCX ingestion using python-docx. Groups paragraphs into blocks of ~15
so citations can reference an approximate section, and extracts tables."""
import docx


def load_docx(path: str) -> list[dict]:
    d = docx.Document(path)
    docs = []

    block, block_start = [], 1
    for i, para in enumerate(d.paragraphs, start=1):
        if para.text.strip():
            block.append(para.text.strip())
        if len(block) >= 15:
            docs.append({
                "text": "\n".join(block),
                "location": f"paragraphs {block_start}-{i}",
                "page": None,
            })
            block, block_start = [], i + 1
    if block:
        docs.append({
            "text": "\n".join(block),
            "location": f"paragraphs {block_start}-{len(d.paragraphs)}",
            "page": None,
        })

    for t_idx, table in enumerate(d.tables, start=1):
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        docs.append({
            "text": "\n".join(rows),
            "location": f"table {t_idx}",
            "page": None,
        })

    return docs
